from pptx import Presentation
from PIL import ImageFont
import sys
p=sys.argv[1]
prs=Presentation(p)
font_path='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
cache={}
def font(sz,bold=False):
    path='/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf' if bold else font_path
    key=(path,int(sz*2))
    if key not in cache: cache[key]=ImageFont.truetype(path,max(6,int(sz*96/72)))
    return cache[key]
def wrap_lines(txt,sz,width_in,bold=False):
    f=font(sz,bold);maxw=width_in*96;lines=0
    for raw in txt.split('\n'):
        words=raw.split()
        if not words:lines+=1;continue
        cur=words[0]
        for w in words[1:]:
            test=cur+' '+w
            if f.getlength(test)<=maxw:cur=test
            else:lines+=1;cur=w
        lines+=1
    return lines
issues=[]
for si,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        if not getattr(sh,'has_text_frame',False) or not sh.text.strip():continue
        tf=sh.text_frame
        w=sh.width/914400-(tf.margin_left+tf.margin_right)/914400
        h=sh.height/914400-(tf.margin_top+tf.margin_bottom)/914400
        if w<=0 or h<=0:continue
        need=0
        for pgh in tf.paragraphs:
            if not pgh.text:continue
            sizes=[r.font.size.pt for r in pgh.runs if r.font.size]
            sz=max(sizes) if sizes else (pgh.font.size.pt if pgh.font.size else 18)
            bold=any(r.font.bold for r in pgh.runs)
            need+=wrap_lines(pgh.text,sz,w,bold)*(sz*1.16/72)
            if pgh.space_after:need+=pgh.space_after.pt/72
        ratio=need/h if h else 0
        if ratio>1.03:issues.append((round(ratio,3),si,round(w,2),round(h,2),round(need,2),sh.text[:100].replace('\n',' / ')))
for row in sorted(issues,reverse=True)[:50]:print(row)
print('flagged',len(issues),'of',sum(1 for s in prs.slides for sh in s.shapes if getattr(sh,'has_text_frame',False) and sh.text.strip()))
if issues:sys.exit(2)
